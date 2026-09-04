from pathlib import Path

from openpyxl import load_workbook

from apps.studio.config import get_default_config
from apps.studio.plan_engine import construir_nombre_iteracion, construir_plan_medida, generar_plan_estudio
from core.postprocess.data import guardar_resumen_studio_excel
from apps.stress.plan_engine import construir_plan_medida as construir_plan_stress
from core.measure.run_ngu401 import construir_plan_medida as construir_plan_ngu401
from core.instrument.solar_simulator import etiquetar_canal_ossila


def _build_cfg():
    cfg = get_default_config()
    cfg["simulador_solar_activo"] = True
    cfg["eje_estructura_activo"] = True
    cfg["barrido_motor_activo"] = False
    cfg["irradiancia_modo"] = "potencia"
    cfg["irradiancia_potencia"] = {
        "p_inicial_mW_cm2": 10.0,
        "p_final_mW_cm2": 20.0,
        "paso_mW_cm2": 10.0,
        "espera_estabilizacion_s": 1.0,
        "espera_encendido_medida_s": 1.0,
    }
    cfg["estructura"] = {
        "estructuras": ["Estructura 1", "Estructura 2"],
        "espera_estabilizacion_s": 0.5,
        "keithley_por_estructura": {
            "Estructura 1": {"recurso_visa": "AUTO", "modo_medida": "completa"},
            "Estructura 2": {"recurso_visa": "AUTO", "modo_medida": "completa"},
        },
    }
    cfg["irradiancia_combinacion"] = {
        "canales_combinacion": [{"canal": "950", "intensidades": "100"}],
        "relaciones": [],
        "espera_estabilizacion_s": 1.0,
        "espera_encendido_medida_s": 1.0,
        "cada_n_medidas_estructura": 1,
        "tiempo_enfriado_s": 30.0,
    }
    cfg["irradiancia_multiples_combinaciones"] = {
        "combinaciones": [],
        "espera_estabilizacion_s": 1.0,
        "espera_encendido_medida_s": 1.0,
        "cada_n_medidas_estructura": 1,
        "tiempo_enfriado_s": 30.0,
    }
    return cfg


def test_generar_plan_estudio_ordena_medidas_por_estructura_y_cooldown():
    cfg = _build_cfg()

    plan = generar_plan_estudio(cfg)

    assert plan[0]["tipo"] == "medida"
    assert plan[0]["estructura"] == "Estructura 1"
    assert plan[1]["tipo"] == "enfriar"
    assert plan[1]["duracion_s"] == 30.0
    assert plan[2]["tipo"] == "medida"
    assert plan[2]["estructura"] == "Estructura 2"


def test_construir_nombre_iteracion_omite_ejes_inactivos():
    nombre = construir_nombre_iteracion(
        "medida studio",
        {
            "estructura": "Estructura 1",
            "solar_modo": "off",
            "solar_params": {},
            "motor_activo": False,
        },
    )

    assert nombre == "medida_studio__Estructura_1"


def test_construir_nombre_iteracion_usa_led_y_unidades_fisicas():
    nombre = construir_nombre_iteracion(
        "medida",
        {
            "estructura": "Estructura 1",
            "solar_modo": "combinacion",
            "solar_params": {"canales": ["950", "660"], "intensidades": [100, 50]},
            "motor_activo": True,
            "eje": "inclinacion",
            "posicion_fisica": 15.0,
            "unidad_fisica": "deg",
        },
    )

    assert nombre == "medida__Estructura_1__950nm_100pct-660nm_050pct__inclinacion_15.000deg"


def test_etiquetar_canal_ossila_usa_longitud_de_onda():
    assert etiquetar_canal_ossila("11") == "950"
    assert etiquetar_canal_ossila("8") == "660"
    assert etiquetar_canal_ossila("cool_white") == "cool_white"


def test_construir_nombre_iteracion_no_expone_canal_interno():
    nombre = construir_nombre_iteracion(
        "medida",
        {
            "solar_modo": "longitud_onda",
            "solar_params": {"canal": "11", "intensidad_pct": 50},
            "motor_activo": False,
        },
    )

    assert nombre == "medida__950nm_050pct"


def test_guardar_resumen_studio_excel_es_una_hoja_sin_rutas(tmp_path):
    cfg = get_default_config()
    cfg["carpeta_salida"] = str(tmp_path)
    cfg["nombre_carpeta_medida"] = "medida"
    cfg["nombre_medida"] = "studio"

    ruta = guardar_resumen_studio_excel(cfg, [
        {
            "Iteración": 1,
            "Nombre iteración": "studio__Estructura_1",
            "Estructura": "Estructura 1",
            "LED 950nm (%)": 100,
            "Estado": "correcta",
            "Puntos": 3,
            "Archivo Excel": "medida.xlsx",
            "Archivo PNG": "medida.png",
            "Voc (V)": 0.5,
        }
    ])

    libro = load_workbook(Path(ruta), read_only=True)
    assert libro.sheetnames == ["resumen_studio"]
    encabezados = list(next(libro["resumen_studio"].iter_rows(values_only=True)))
    assert "LED 950nm (%)" in encabezados
    assert any(str(columna).startswith("Voc (") for columna in encabezados)
    assert "Archivo Excel" not in encabezados
    assert "Archivo PNG" not in encabezados


def test_todos_los_planificadores_iv_recorrer_en_decreciente():
    cfg = {
        "modo_medida": "completa",
        "directa": {"v_inicial_mV": 0.0, "v_final_mV": 100.0, "paso_mV": 25.0},
        "inversa": {"v_inicial_mV": None, "v_final_V": -0.1, "paso_mV": 25.0},
    }

    for construir_plan in (construir_plan_medida, construir_plan_stress, construir_plan_ngu401):
        segmentos = construir_plan(cfg)
        assert all(
            all(tensiones[i] >= tensiones[i + 1] for i in range(len(tensiones) - 1))
            for tensiones in (segmento["tensiones_V"] for segmento in segmentos)
        )
